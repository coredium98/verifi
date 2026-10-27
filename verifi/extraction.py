"""Financial statement extraction (Excel and PDF) and data quality assessment."""

from __future__ import annotations

import io
import math
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any

# Ordered rules — first match wins. More specific patterns go first.
MATCH_RULES: list[tuple[str, list[str], list[str]]] = [
    ("operating_cash_flow", ["operating cash flow", "cash from operat", "cash flow from operat"], []),
    ("ebitda", ["ebitda"], []),
    ("gross_profit", ["gross profit", "gross margin"], ["margin %", "margin%"]),
    ("revenue", ["revenue / sales", "net sales", "total revenue", "revenue"], ["cost", "other revenue"]),
    ("cogs", ["cost of goods", "cogs", "cost of sales", "direct cost"], []),
    ("net_income", ["net income", "net profit", "net earnings", "net loss"], ["change in", "add:"]),
    (
        "owner_comp",
        [
            "owner / officer",
            "owner/officer",
            "officer compensation",
            "officer salary",
            "owner compensation",
            "owner salary",
        ],
        [],
    ),
    ("accounts_receivable", ["accounts receivable", "trade receivable"], ["change in", "change -"]),
    ("depreciation", ["depreciation", "amortization"], ["accumulated", "add:"]),
    ("interest", ["interest expense"], []),
    ("tax", ["income tax", "tax expense"], []),
]

MANUAL_FIELDS: list[tuple[str, str]] = [
    ("revenue", "Revenue"),
    ("cogs", "Cost of Goods Sold"),
    ("gross_profit", "Gross Profit"),
    ("net_income", "Net Income"),
    ("operating_cash_flow", "Operating Cash Flow"),
    ("accounts_receivable", "Accounts Receivable (Balance Sheet)"),
    ("owner_comp", "Owner / Officer Compensation"),
    ("ebitda", "EBITDA"),
]


def _normalize(text: Any) -> str:
    if text is None:
        return ""
    return str(text).lower().strip()


def _parse_number(val: Any) -> float | None:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        if math.isnan(val):
            return None
        return float(val)
    s = (
        str(val)
        .replace(",", "")
        .replace("$", "")
        .replace("(", "-")
        .replace(")", "")
        .replace("%", "")
        .strip()
    )
    if not s or s.startswith("#") or s.startswith("="):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _classify(label: Any) -> str | None:
    t = _normalize(label)
    if not t:
        return None
    for key, required, excluded in MATCH_RULES:
        if any(r in t for r in required):
            if not any(e in t for e in excluded):
                return key
    return None


def _better(existing: list[float], new: list[float]) -> list[float]:
    """Return whichever list has more non-zero, non-None values."""

    def score(lst: list[float]) -> int:
        return sum(1 for v in lst if v is not None and v != 0)

    return new if score(new) > score(existing) else existing


def _ingest_rows(rows: list[list[Any]]) -> dict[str, list[float]]:
    data: dict[str, list[float]] = {}
    for row in rows:
        if not row or row[0] is None:
            continue
        key = _classify(row[0])
        if key is None:
            continue
        nums: list[float] = []
        for cell in row[1:]:
            if isinstance(cell, str) and (cell.startswith("#") or cell.startswith("=")):
                continue
            n = _parse_number(cell)
            if n is not None:
                nums.append(n)
        nums = nums[:3]
        if not nums:
            continue
        if key in data:
            data[key] = _better(data[key], nums)
        else:
            data[key] = nums
    return data


def _extract_xlsx_stdlib(file_source: bytes | io.BytesIO | str | Path) -> dict[str, list[float]]:
    """Standard-library OpenXML (.xlsx) fallback parser."""
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    if isinstance(file_source, bytes):
        stream: io.BytesIO | str | Path = io.BytesIO(file_source)
    else:
        if hasattr(file_source, "seek"):
            file_source.seek(0)
        stream = file_source

    rows: list[list[Any]] = []
    with zipfile.ZipFile(stream) as z:
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in z.namelist():
            ss_root = ET.fromstring(z.read("xl/sharedStrings.xml"))
            for si in ss_root.findall("s:si", ns):
                parts = [t.text or "" for t in si.findall(".//s:t", ns)]
                shared_strings.append("".join(parts))

        sheet_names = sorted(
            name for name in z.namelist() if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
        )
        if not sheet_names:
            return {}

        sheet = ET.fromstring(z.read(sheet_names[0]))
        for row_el in sheet.findall(".//s:row", ns):
            row_vals: list[Any] = []
            for cell_el in row_el.findall("s:c", ns):
                cell_type = cell_el.attrib.get("t")
                val_el = cell_el.find("s:v", ns)
                inline_el = cell_el.find("s:is/s:t", ns)
                if cell_type == "s" and val_el is not None and val_el.text is not None:
                    idx = int(val_el.text)
                    row_vals.append(shared_strings[idx] if 0 <= idx < len(shared_strings) else "")
                elif inline_el is not None:
                    row_vals.append(inline_el.text)
                elif val_el is not None:
                    row_vals.append(val_el.text)
                else:
                    row_vals.append(None)
            if row_vals:
                rows.append(row_vals)

    return _ingest_rows(rows)


def extract_from_excel(file_bytes: bytes | io.BytesIO | str | Path) -> dict[str, list[float]]:
    """Extract key financial line items from an Excel workbook (.xlsx / .xls)."""
    try:
        from openpyxl import load_workbook  # type: ignore[import-not-found]

        if hasattr(file_bytes, "seek"):
            file_bytes.seek(0)
        wb = load_workbook(file_bytes, data_only=True)
        ws = wb.active
        rows = [[cell.value for cell in row] for row in ws.iter_rows()]
        return _ingest_rows(rows)
    except Exception:
        pass

    try:
        import pandas as pd  # type: ignore[import-not-found]

        if hasattr(file_bytes, "seek"):
            file_bytes.seek(0)
        df = pd.read_excel(file_bytes, header=None)
        rows = [list(row) for _, row in df.iterrows()]
        return _ingest_rows(rows)
    except Exception:
        pass

    try:
        return _extract_xlsx_stdlib(file_bytes)
    except Exception:
        return {}


def extract_from_pdf(file_bytes: bytes | io.BytesIO | str | Path) -> dict[str, list[float]]:
    """Extract key financial line items from PDF tables using pdfplumber."""
    import pdfplumber  # type: ignore[import-not-found]

    rows: list[list[Any]] = []
    if hasattr(file_bytes, "seek"):
        file_bytes.seek(0)
    with pdfplumber.open(file_bytes) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                for row in table:
                    if row:
                        rows.append(list(row))
    return _ingest_rows(rows)


def _parse_years_claimed(years_claimed: str | int | None) -> int:
    if isinstance(years_claimed, int):
        return max(1, years_claimed)
    if isinstance(years_claimed, str):
        token = years_claimed.strip().split()[0] if years_claimed.strip() else "1"
        try:
            return max(1, int(token))
        except ValueError:
            return 1
    return 1


def assess_quality(
    data: dict[str, list[float]],
    years_claimed: str | int = "3 years",
) -> dict[str, Any]:
    """Assess completeness and multi-year depth of extracted financial statement data."""
    has_revenue = "revenue" in data
    has_gp = "gross_profit" in data or ("cogs" in data and has_revenue)
    has_ni = "net_income" in data
    has_income = has_revenue and has_gp and has_ni
    has_balance = "accounts_receivable" in data
    has_cash_flow = "operating_cash_flow" in data

    expected_years = _parse_years_claimed(years_claimed)
    years_found = max((len(v) for v in data.values() if isinstance(v, list)), default=0)

    core_count = sum([has_revenue, has_gp, has_ni, has_balance, has_cash_flow])

    if core_count >= 4:
        score = "Sufficient"
    elif core_count >= 2:
        score = "Partial"
    else:
        score = "Insufficient"

    missing: list[str] = []
    if not has_income:
        missing.append("complete income statement (revenue, gross profit, net income)")
    if not has_balance:
        missing.append("balance sheet with accounts receivable")
    if not has_cash_flow:
        missing.append("cash flow statement or operating cash flow figure")
    if "owner_comp" not in data:
        missing.append("owner / officer compensation detail")

    if years_found < expected_years and core_count > 0:
        missing.append(
            f"historical columns matching selection ({years_found} of {expected_years} years found)"
        )
        if expected_years >= 2 and years_found < 2 and score == "Sufficient":
            score = "Partial"

    return {
        "has_income_stmt": has_income,
        "has_balance": has_balance,
        "has_cash_flow": has_cash_flow,
        "score": score,
        "missing": missing,
        "found_keys": list(data.keys()),
        "years_claimed": expected_years,
        "years_found": years_found,
    }
